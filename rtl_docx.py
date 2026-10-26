#!/usr/bin/env python3
"""Mark a .docx as right-to-left so Hebrew paragraphs align from the right.

Word mirrors left and right once w:bidi is set, so paragraph alignment is
stored as left. That is what shows up on the right edge.

Run:
  python rtl_docx.py price_action.docx
  python rtl_docx.py price_action.docx --out price_action_rtl.docx

The input file is never overwritten. With no --out, the result is <name>_rtl.docx
beside the input.
"""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_PARAGRAPH_STYLES = ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4")


def _set_on(parent, tag: str) -> None:
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    element.set(qn("w:val"), "1")


def _set_jc_left(pPr) -> None:
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), "left")


def _rtl_paragraph(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    _set_on(pPr, "w:bidi")
    # Word mirrors left/right when w:bidi is set, so jc=left sits on the right edge.
    _set_jc_left(pPr)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        _set_on(run._r.get_or_add_rPr(), "w:rtl")


def _rtl_container(paragraphs, tables) -> None:
    for paragraph in paragraphs:
        _rtl_paragraph(paragraph)
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                _rtl_container(cell.paragraphs, cell.tables)


def _rtl_styles(doc: Document) -> None:
    for name in _PARAGRAPH_STYLES:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        pPr = style.element.get_or_add_pPr()
        _set_on(pPr, "w:bidi")
        _set_jc_left(pPr)


def _rtl_sections(doc: Document) -> None:
    for section in doc.sections:
        sectPr = section._sectPr
        _set_on(sectPr, "w:bidi")
        if sectPr.find(qn("w:headerReference")) is not None:
            _rtl_container(section.header.paragraphs, section.header.tables)
        if sectPr.find(qn("w:footerReference")) is not None:
            _rtl_container(section.footer.paragraphs, section.footer.tables)


def _rtl_theme_lang(doc: Document) -> None:
    settings = doc.settings.element
    theme = settings.find(qn("w:themeFontLang"))
    if theme is None:
        theme = OxmlElement("w:themeFontLang")
        settings.append(theme)
    if not theme.get(qn("w:val")):
        theme.set(qn("w:val"), "en-US")
    theme.set(qn("w:bidi"), "he-IL")


def apply_rtl(doc: Document) -> None:
    """Mark styles, sections, and every paragraph in doc as right-to-left."""
    _rtl_styles(doc)
    _rtl_sections(doc)
    _rtl_container(doc.paragraphs, doc.tables)
    _rtl_theme_lang(doc)


def _default_out(src: Path) -> Path:
    return src.with_name(f"{src.stem}_rtl{src.suffix}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert an existing .docx to right-to-left layout.")
    parser.add_argument("docx", help="Existing .docx to convert. This file is not modified.")
    parser.add_argument("--out", default=None, help="Output .docx. Defaults to <name>_rtl.docx next to the input.")
    args = parser.parse_args()

    src = Path(args.docx)
    if not src.is_file():
        raise SystemExit(f"Not a file: {src}")
    dest = Path(args.out) if args.out else _default_out(src)
    if dest.resolve() == src.resolve():
        parser.error("--out must be a different file; the input is not modified")
    doc = Document(src)
    apply_rtl(doc)
    doc.save(dest)
    print(f"Saved right-to-left document: {dest}")


if __name__ == "__main__":
    main()
